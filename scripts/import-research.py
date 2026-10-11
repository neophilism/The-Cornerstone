"""Reproduce public research projections from the approved, unmodified originals.

Requires Python 3.11+, openpyxl and python-docx for this offline import only.
Never fetches citations, infers certification, geocodes a person, or edits a source.
"""
from pathlib import Path
import hashlib
import json
import re
import sys
from datetime import date, datetime

import openpyxl
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
EXPECTED = {
    "Event evidence": 231, "Citation audit": 231, "Workforce evidence": 128,
    "Worksites": 120, "Method": 17, "Coverage audit": 37,
}


def value(v):
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    return str(v)


def extract():
    manifest = json.loads((ROOT / "policy/SOURCE_MANIFEST.json").read_text())
    for source in manifest["sources"]:
        data = (ROOT / source["path"]).read_bytes()
        if len(data) != source["bytes"] or hashlib.sha256(data).hexdigest() != source["sha256"]:
            raise ValueError(f"Source changed without a new reviewed version: {source['path']}")
    workbook = openpyxl.load_workbook(ROOT / "policy/Event_Evidence.xlsx", read_only=True, data_only=True)
    formula_workbook = openpyxl.load_workbook(ROOT / "policy/Event_Evidence.xlsx", read_only=True, data_only=False)
    if list(EXPECTED) != workbook.sheetnames:
        raise ValueError("Workbook sheet inventory changed")
    # These supplied sheets contain source records, not business formulas. A future
    # formula-bearing source needs an explicit cached-value/recalculation policy.
    for sheet in formula_workbook:
        if any(c.data_type == "f" for row in sheet for c in row):
            raise ValueError(f"Formula requires reviewed extraction policy: {sheet.title}")
    sheets = []
    for sheet in workbook:
        rows = list(sheet.values)
        headers = [value(v) for v in rows[3]]
        if any(not h for h in headers) or len(set(headers)) != len(headers):
            raise ValueError(f"Invalid headers: {sheet.title}")
        records = [
            {"row": n, "values": [value(v) for v in row[:len(headers)]]}
            for n, row in enumerate(rows[4:], start=5) if any(v is not None for v in row)
        ]
        if len(records) != EXPECTED[sheet.title]:
            raise ValueError(f"Row reconciliation failed: {sheet.title}")
        sheets.append({"name": sheet.title, "heading": value(rows[0][0]),
                       "context": value(rows[1][0]), "headerRow": 4,
                       "headers": headers, "records": records})
    ids = [r["values"][0] for r in sheets[0]["records"]]
    if len(set(ids)) != len(ids) or any(not re.fullmatch(r"C-\d{3}[A-Z]?", i) for i in ids):
        raise ValueError("Duplicate or invalid research event IDs")
    document = Document(ROOT / "policy/Displacement_Event_Census.docx")
    census = []
    active = None
    introduction = []
    for number, p in enumerate(document.paragraphs, start=1):
        text = p.text.strip()
        match = re.match(r"^(C-\d{3}[A-Z]?)\s+(.+)$", text)
        if match:
            active = {"id": match[1], "title": match[2], "paragraph": number, "paragraphs": []}
            census.append(active)
        target = active["paragraphs"] if active is not None else introduction
        if text:
            links = []
            for h in p._p.iter(qn("w:hyperlink")):
                relationship = document.part.rels.get(h.get(qn("r:id")))
                if relationship is not None:
                    links.append({"text": "".join(t.text or "" for t in h.iter(qn("w:t"))), "url": relationship.target_ref})
            target.append({"paragraph": number, "text": text, "links": links})
    census_ids = [r["id"] for r in census]
    if set(ids) != set(census_ids) or len(census_ids) != len(ids):
        raise ValueError("Census/workbook event IDs do not reconcile")
    for sheet in sheets[1:4]:
        column = 0 if sheet["name"] == "Citation audit" else 1
        for row in sheet["records"]:
            linked = re.findall(r"C-\d{3}[A-Z]?", row["values"][column])
            if not linked or not set(linked).issubset(ids):
                raise ValueError(f"Unresolved event link: {sheet['name']} row {row['row']}")
    sources = [{k: v for k, v in source.items() if k != "library_file_id"} for source in manifest["sources"]]
    return {"version": 1, "classification": "research-only", "publicationApproved": "2026-10-11",
            "workbookReviewed": "2026-09-28", "sources": sources, "sheets": sheets,
            "censusIntroduction": introduction, "census": census}


def main():
    data = extract()
    output = ROOT / "packages/content/src/research-data.json"
    contents = json.dumps(data, ensure_ascii=False, indent=2) + "\n"
    if "--check" in sys.argv:
        if json.loads(output.read_text()) != data:
            raise ValueError("Generated research projection differs; rerun importer and review changes")
    else:
        output.write_text(contents)
    report = {"researchEvents": len(data["census"]), "censusIDsReconciled": True,
              "sheets": {s["name"]: len(s["records"]) for s in data["sheets"]},
              "sourceHashesVerified": len(data["sources"]), "certificationsCreated": 0,
              "citationsReverifiedByImport": False, "geocodingPerformed": False}
    if "--check" not in sys.argv:
        (ROOT / "docs/RESEARCH_RECONCILIATION.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report))


if __name__ == "__main__":
    main()
